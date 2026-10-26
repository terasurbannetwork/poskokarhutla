"""PDF pratinjau deck tanpa LibreOffice/PowerPoint: tiap shape pptx (kotak, oval, teks, gambar, garis) ditulis ulang sebagai HTML
berposisi mutlak, lalu dicetak Chrome headless. Hanya mendukung bentuk yang dipakai buat.py.
Pakai: python3 pratinjau.py deck.pptx keluar.pdf <chrome-headless-shell>"""
import base64, html, os, subprocess, sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
SRC, OUT, CH = sys.argv[1:4]
E = 12700
prs = Presentation(SRC); W, H = prs.slide_width / E, prs.slide_height / E
def pt(v): return v / E
def warna(fill):
    try:
        if fill.type == 1: return '#' + str(fill.fore_color.rgb)
    except Exception: pass
    return 'transparent'
el = []
for sl in prs.slides:
    el.append(f"<section style='width:{W}pt;height:{H}pt'>")
    for sh in sl.shapes:
        x, y, w, h = pt(sh.left), pt(sh.top), pt(sh.width), pt(sh.height)
        st = f'left:{x:.2f}pt;top:{y:.2f}pt;width:{w:.2f}pt;height:{h:.2f}pt;'
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            b64 = base64.b64encode(sh.image.blob).decode()
            el.append(f"<img style='{st}' src='data:{sh.image.content_type};base64,{b64}'>"); continue
        if sh.shape_type == MSO_SHAPE_TYPE.LINE or sh.shape_type is None and sh.__class__.__name__ == 'Connector':
            ln = sh.line; c = '#' + str(ln.color.rgb); t = (ln.width or E) / E
            x1, y1, x2, y2 = pt(sh.begin_x), pt(sh.begin_y), pt(sh.end_x), pt(sh.end_y)
            if abs(y1 - y2) < .01: el.append(f"<div style='left:{min(x1, x2)}pt;top:{y1 - t / 2}pt;width:{abs(x2 - x1)}pt;height:0;border-top:{t}pt solid {c}'></div>")
            else: el.append(f"<div style='left:{x1 - t / 2}pt;top:{min(y1, y2)}pt;width:0;height:{abs(y2 - y1)}pt;border-left:{t}pt solid {c}'></div>")
            continue
        css = st
        if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            css += f'background:{warna(sh.fill)};'
            ln = sh.line
            try:
                if ln.fill.type == 1:
                    dash = 'dashed' if ln.dash_style else 'solid'; css += f'outline:{(ln.width or E) / E}pt {dash} #{ln.color.rgb};outline-offset:-{(ln.width or E) / E / 2}pt;'
            except Exception: pass
            if sh.auto_shape_type == MSO_SHAPE.OVAL: css += 'border-radius:50%;'
            elif sh.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE: css += f'border-radius:{sh.adjustments[0] * min(w, h):.2f}pt;'
        inner = ''
        if sh.has_text_frame and sh.text_frame.text:
            tf = sh.text_frame; p = tf.paragraphs[0]
            al = {PP_ALIGN.CENTER: 'center', PP_ALIGN.RIGHT: 'flex-end', PP_ALIGN.JUSTIFY: 'flex-start'}.get(p.alignment, 'flex-start')
            ta = {PP_ALIGN.CENTER: 'center', PP_ALIGN.RIGHT: 'right', PP_ALIGN.JUSTIFY: 'justify'}.get(p.alignment, 'left')
            va = {MSO_ANCHOR.TOP: 'flex-start', MSO_ANCHOR.BOTTOM: 'flex-end'}.get(tf.vertical_anchor, 'center')
            spans = ''
            for r in p.runs:
                f = r.font; spc = r._r.find('{http://schemas.openxmlformats.org/drawingml/2006/main}rPr')
                ls = int(spc.get('spc', 0)) / 100 if spc is not None else 0
                spans += f"<span style='font-size:{f.size.pt}pt;font-weight:{700 if f.bold else 400};color:#{f.color.rgb};letter-spacing:{ls}pt'>{html.escape(r.text)}</span>"
            lh = p.line_spacing if isinstance(p.line_spacing, float) else 1.17
            ws = 'normal' if tf.word_wrap else 'nowrap'
            fs = max((r.font.size.pt for r in p.runs if r.font.size), default=8)
            inner = f"<div style='display:flex;height:100%;align-items:{va};justify-content:{al}'><div style='text-align:{ta};line-height:{lh};white-space:{ws};width:100%;font-size:{fs}pt'>{spans}</div></div>"
        el.append(f"<div style='{css}'>{inner}</div>")
    el.append('</section>')
doc = f"""<!doctype html><html><head><meta charset='utf-8'>
<link href='https://fonts.googleapis.com/css2?family=Roboto:wght@400;700&display=block' rel='stylesheet'>
<style>@page{{size:{W}pt {H}pt;margin:0}} body{{margin:0;font-family:Roboto,sans-serif}} section{{position:relative;overflow:hidden;page-break-after:always;background:#fff}}
section>*{{position:absolute;box-sizing:border-box}}</style></head><body>{''.join(el)}</body></html>"""
open(OUT + '.html', 'w').write(doc)
subprocess.run([CH, '--headless', '--no-pdf-header-footer', f'--print-to-pdf={os.path.abspath(OUT)}', '--virtual-time-budget=8000',
                'file://' + os.path.abspath(OUT + '.html')], check=True, capture_output=True)
os.remove(OUT + '.html'); print('pdf', OUT)
