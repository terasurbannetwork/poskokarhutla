"""Paparan Karhutla Provinsi Jambi, format "Paparan Karhutla Provinsi Lampung 241026" (4 slide: provinsi, kabupaten teratas,
Kawasan Berbak, kebutuhan pemadaman). Deck dibangun dari nol dengan python-pptx; semua angka dari data.json (data.py), peta dari peta.py.
Pakai: python3 buat.py keluar.pptx"""
import datetime as dt, json, math, re, sys
from collections import Counter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Pt

D = json.load(open('data.json')); OUT = sys.argv[1]
NAVY, ORG, MERAH, BIRU, ABU, SLATE, TEKS = '051C2C', 'D9600B', 'C8102E', '2F86D1', '6B7280', '3F5F80', '1F2937'
FONT = 'Roboto'
BLN = ['Jan', 'Feb', 'Mar', 'Apr', 'Mei', 'Jun', 'Jul', 'Agu', 'Sep', 'Okt', 'Nov', 'Des']
BLN_P = ['Januari', 'Februari', 'Maret', 'April', 'Mei', 'Juni', 'Juli', 'Agustus', 'September', 'Oktober', 'November', 'Desember']
f0 = lambda v: f'{int(round(v)):,}'.replace(',', '.')
f1 = lambda v: (f'{v:.1f}'.replace('.', ',') if round(v, 1) != int(round(v, 1)) else f0(v))
tg = lambda d: f"{int(d[8:])} {BLN[int(d[5:7]) - 1]}"          # '2026-10-25' -> '25 Okt'
CUT = dt.datetime.fromisoformat(D['cut']); HARI = D['hari']; SEM = D['sementara']
JAM = CUT.strftime('%H.%M')
CUT_TXT = f"{CUT.day} {BLN[CUT.month - 1]} {CUT.year}, {JAM} WIB"
HD = tg(HARI)                                                     # '25 Okt'
HD_S = f"{HD} s.d. {JAM}" if SEM else HD
D8 = D['D8']; TGL = [str(int(d[8:])) for d in D8]
CATATAN = '*Data berdasarkan Geographic Information System (GIS) dengan satelit NASA (FIRMS: VIIRS dan MODIS).'
if not SEM and CUT.date().isoformat() != HARI: CATATAN += f' Angka per {HD} (hari penuh terakhir).'

prs = Presentation(); prs.slide_width, prs.slide_height = Pt(960), Pt(540)
BL = prs.slide_layouts[6]

def rgb(c): return RGBColor.from_string(c)
class Sl:
    def __init__(s, nomor):
        s.sl = prs.slides.add_slide(BL); s.sh = s.sl.shapes; s.nomor = nomor
    def kotak(s, x, y, w, h, isi=None, garis=None, r=0, bentuk=None, tebal=.75, putus=False):
        b = bentuk or (MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE)
        sp = s.sh.add_shape(b, Pt(x), Pt(y), Pt(w), Pt(h))
        if r and b == MSO_SHAPE.ROUNDED_RECTANGLE: sp.adjustments[0] = min(.5, r / min(w, h))
        if isi: sp.fill.solid(); sp.fill.fore_color.rgb = rgb(isi)
        else: sp.fill.background()
        if garis: sp.line.color.rgb = rgb(garis); sp.line.width = Pt(tebal)
        else: sp.line.fill.background()
        if putus: sp.line.dash_style = MSO_LINE.DASH
        sp.shadow.inherit = False
        sp.text_frame.text = ''
        return sp
    def bulat(s, x, y, d, isi, garis=None): return s.kotak(x, y, d, d, isi, garis, bentuk=MSO_SHAPE.OVAL)
    def teks(s, x, y, w, h, runs, sz=8, warna=TEKS, tebal=False, rata='l', v='m', spasi=None, wrap=True):
        """runs: str atau [(teks, warna|None, tebal|None, ukuran|None)]"""
        tb = s.sh.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h)); tf = tb.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.word_wrap = wrap
        tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[v]
        p = tf.paragraphs[0]; p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT, 'j': PP_ALIGN.JUSTIFY}[rata]
        if spasi: p.line_spacing = spasi
        for t in ([(runs,)] if isinstance(runs, str) else runs):
            t = list(t) + [None] * (4 - len(t)); r = p.add_run(); r.text = t[0]; f = r.font
            f.name = FONT; f.size = Pt(t[3] or sz); f.bold = tebal if t[2] is None else t[2]; f.color.rgb = rgb(t[1] or warna)
            if t[0].isupper() and len(t[0]) > 5 and (t[3] or sz) <= 8 and not any(ch.isdigit() for ch in t[0]): r._r.get_or_add_rPr().set('spc', '120' if (t[3] or sz) > 5.5 else '40')
        return tb
    def gambar(s, path, x, y, w=None, h=None): return s.sh.add_picture(path, Pt(x), Pt(y), Pt(w) if w else None, Pt(h) if h else None)
    def ikon(s, nama, warna, x, y, d): return s.gambar(f'ik/{nama}-{warna}.png', x, y, d, d)
    def garis(s, x1, y1, x2, y2, warna='D5DBE3', tebal=.75):
        c = s.sh.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2)); c.line.color.rgb = rgb(warna); c.line.width = Pt(tebal); return c

    # ---------------------------------------------------------------- komponen
    def kepala(s, judul, sorot, ekor, cutoff=True, dua_baris=False):
        h = 64 if dua_baris else 40
        n = max(1, len(judul + (sorot or '') + ekor)); sz = min(22, 646 / (n * .485))
        if isinstance(sorot, str): s.teks(43, 24, 650 if cutoff else 820, h, [(judul, NAVY), (sorot, ORG), (ekor, NAVY)], sz, tebal=True, v='t' if dua_baris else 'm', spasi=1.0, wrap=False)
        if cutoff:
            s.kotak(696, 32, 186, 17, NAVY, r=8.5); s.ikon('kalender', 'putih', 704, 35.5, 10)
            s.teks(719, 32, 70, 17, 'CUT OFF DATA', 5.8, 'C9D1E3', True)
            s.teks(780, 32, 96, 17, CUT_TXT, 7.6, 'FFFFFF', True, 'r')
        s.gambar('logo_bnpb.png', 889.4, 28.8, 27.4, 34.6)
        yl = 84 if dua_baris else 71
        s.garis(43, yl, 917, yl, '1F2937', .9)
        s.garis(43, 508, 917, 508, '1F2937', .9)
        s.teks(43, 512, 432, 15.8, 'Badan Nasional Penanggulangan Bencana  |  Copyright 2026', 8.2, '1F2937')
        s.teks(873.6, 512, 43.2, 15.8, str(s.nomor), 8.2, '1F2937', True, 'r')
        return yl
    def catatan(s, teks):
        s.kotak(43, 490, 2.6, 12, BIRU); s.teks(50, 489, 860, 14, teks, 6.6, '1F2937')
    def label(s, x, y, teks, isi=NAVY, ik='api', warna='FFFFFF', w=None):
        w = w or 4.9 * len(teks) + 30; s.kotak(x, y, w, 14, isi, r=7); s.ikon(ik, 'putih', x + 7, y + 3, 8); s.teks(x + 19, y, w - 20, 14, teks, 6.2, warna, True)
    def kartu(s, x, y, w, h, angka, ket, ik, warna_ik, tint, warna_angka=NAVY, sz=17, ket_sz=7.4):
        s.kotak(x, y, w, h, 'FFFFFF', 'D5DBE3', r=6)
        d = 24; s.bulat(x + 9, y + (h - d) / 2, d, tint); s.ikon(ik, warna_ik, x + 9 + 5, y + (h - d) / 2 + 5, 14)
        s.teks(x + 41, y + 2, w - 46, h * .5, angka, sz, warna_angka, True, v='b')
        s.teks(x + 41, y + h * .53, w - 46, h * .45, ket, ket_sz, '374151', v='t', spasi=.95)
    def pil(s, x, y, teks, isi, warna, ik=None, sz=5.9, h=11, k=3.35, c0=10):
        w = k * len(teks) + c0 + (9 if ik else 0)
        s.kotak(x, y, w, h, isi, r=h / 2)
        if ik: s.ikon(ik, ik_warna(warna), x + 4, y + (h - 6.5) / 2, 6.5)
        s.teks(x + (12 if ik else 4), y, w - (14 if ik else 8), h, teks, sz, warna, True, 'l' if ik else 'c')
        return w
    def panel(s, x, y, w, h, judul, ik, warna_lingkar, tint, kanan=None, warna_kanan=None, hh=38):
        s.kotak(x, y, w, h, 'FFFFFF', 'D5DBE3', r=6)
        s.kotak(x + .5, y + .5, w - 1, hh, tint, r=6); s.kotak(x + .5, y + hh - 8, w - 1, 8.5, tint)
        s.garis(x, y + hh, x + w, y + hh, 'D5DBE3')
        s.bulat(x + 10, y + (hh - 23) / 2, 23, warna_lingkar); s.ikon(ik, 'putih', x + 15, y + (hh - 23) / 2 + 5, 13)
        s.teks(x + 40, y, w - 90, hh, judul, 11, NAVY, True)
        if kanan is not None: s.teks(x + w - 60, y, 50, hh, kanan, 15, warna_kanan or NAVY, True, 'r')

def ik_warna(hexc): return {NAVY: 'navy', ORG: 'oranye', MERAH: 'merah', ABU: 'abu', SLATE: 'slate', BIRU: 'biru', '8A6A2E': 'coklat', 'FFFFFF': 'putih'}.get(hexc, 'navy')
CHIP = {'n': ('E6EAF0', NAVY), 'o': ('FBE7D6', ORG), 'm': ('F8DDE0', MERAH), 'g': ('E6EAF0', ABU), 'y': ('F4ECDC', '8A6A2E'), 'b': ('DCEBF8', BIRU)}

def pct(a, b):
    if a == b: return '0%', ABU
    if a == 0: return 'baru', 'A9B4C2'
    v = (b - a) / a * 100; t = f0(abs(v)) + '%'
    return ('▲' + t, ORG) if v > 0 else ('▼' + t, BIRU)
def warna_batang(vals, n, akhir_sementara):
    if n == 0: return 'DCE1E7', None
    d = (vals[n] > vals[n - 1]) - (vals[n] < vals[n - 1])
    if n == len(vals) - 1 and akhir_sementara: return ('FBE1CC' if d > 0 else 'DCEBF8' if d < 0 else 'E6EAF0'), ('F0A35E' if d > 0 else '6FA8DC' if d < 0 else 'A9B4C2')
    pekat = n >= len(vals) - 3
    if d == 0: return 'DCE1E7', None
    return ((ORG if pekat else 'F4C3A0') if d > 0 else (BIRU if pekat else 'B7D3EE')), None

def batang(s, vals, x0, lebar, base, maxh, bw, tgl, pcts=True, sz_v=8, sz_t=6.2, label_akhir=None):
    m = max(vals) or 1; cw = lebar / len(vals)
    for n, v in enumerate(vals):
        cx = x0 + cw * n + cw / 2; h = max(1.6, maxh * v / m); isi, putus = warna_batang(vals, n, SEM)
        s.kotak(cx - bw / 2, base - h, bw, h, isi, putus, r=2.5 if h > 6 else 0, putus=bool(putus), tebal=.9)
        s.teks(cx - 25, base - h - 13, 50, 12, f0(v), sz_v, NAVY, True, 'c', 'b')
        t = tgl[n] if not (n == len(vals) - 1 and label_akhir) else label_akhir
        s.teks(cx - 32, base + 2, 64, 8, t, sz_t, ABU, False, 'c', 't')
        if pcts and n > 0:
            tt, cc = pct(vals[n - 1], v); s.teks(cx - 25, base + 9.5, 50, 8, tt, sz_t, cc, True, 'c', 't')
    s.garis(x0, base, x0 + lebar, base, 'C9D1E3', .6)

def ringkas_ubah(a, b):
    t, _ = pct(a, b)
    return 'tetap' if t == '0%' else 'baru' if t == 'baru' else ('naik ' if t[0] == '▲' else 'turun ') + t[1:]

PV = D['prov']; KAB = D['kab']; K7 = D['k7']; NLV = D['nlv']; KW = D['kaw']; KK = D['kawk']; JK = D['kompleks']
OPP = D['op_prov']; ADA = OPP['ada']; PM = D['pm']; BUT = D['butuh']; KEND = D['kendala']; OPK = D['op_kaw']
hari_ada = ADA['sumber'].split()[-1]                      # tanggal rekap BNPB harian
def kat_pm(v):
    b = PM['batas']; return 'baik' if v <= b[0] else 'sedang' if v <= b[1] else 'tidak sehat' if v <= b[2] else 'sangat tidak sehat' if v <= b[3] else 'berbahaya'
pm_kab = max(PM['kab'], key=lambda r: r[2]) if PM['kab'] else None
pm_v, pm_lok = (pm_kab[2], pm_kab[1]) if pm_kab else (PM['prov'][5], 'rata-rata Jambi')
pm_short = pm_lok.replace('Tanjung Jabung', 'Tanjab'); PMH = tg(PM['hari'])
KAW_JK = [k for k in JK if k['kaw']]
TOPK = max(KAW_JK, key=lambda k: (k['level'], k['luas'])) if KAW_JK else JK[0]
NAMA_LV = ['rutin', 'patroli', 'siaga', 'prioritas']
KAWN = {'tn': 'TN Berbak', 'thr': 'Tahura OKH'}

# =============================================================== slide 1: Provinsi Jambi
s = Sl(1)
s.kepala('Sebaran Titik Panas dan Kebakaran di ', 'Jambi', f" per {int(HARI[8:])} {BLN_P[int(HARI[5:7]) - 1]}")
h, l, b = PV['hot'], PV['luas'], PV['belum']
s.kotak(43, 81, 565, 78, 'EEF1F5', r=8); s.label(52, 87, 'DATA SATELIT')
s.kartu(52, 107, 177, 43, f0(h[-1]), f'titik panas {HD}', 'api', 'oranye', 'FBE7D6', ORG)
s.kartu(237, 107, 177, 43, f0(l[-1]) + ' ha', f'luas terbakar {HD}', 'grid', 'navy', 'E6EAF0', NAVY)
s.kartu(422, 107, 178, 43, f0(b[-1]) + ' ha', f'belum padam {HD}', 'awas', 'merah', 'F8DDE0', MERAH)
s.kotak(630, 81, 287, 78, 'FBEDEF', r=8); s.label(639, 87, 'PERLU PERHATIAN', MERAH, 'awas')
LVK = 'Prioritas' if NLV['Prioritas'] else 'Siaga'
s.kartu(638, 107, 132, 43, f0(NLV[LVK]), f'kompleks api {LVK.lower()}', 'api', 'merah', 'F8DDE0', MERAH)
s.kartu(777, 107, 132, 43, f'PM2.5 {f0(pm_v)}', f'{kat_pm(pm_v)}, {pm_short}, prakiraan {PMH}', 'asap', 'slate', 'E6EAF0', NAVY, sz=14, ket_sz=6.2)
# peta + penanda
P1 = json.load(open('peta1.json')); mw = 340; mh = mw * P1['h'] / P1['w']
s.gambar('peta1.png', 43, 172, mw, mh)
y = 172 + mh + 4; x = 45
for n, (kec, kab, c) in enumerate(D['kec24'][:3], 1):
    t = f"{kec}, {kab.replace('Tanjung Jabung', 'Tanjab')}  "
    w = 3.45 * len(t) + 3.2 * len(f'{c} titik') + 22
    s.kotak(x, y, w, 12, 'FFFFFF', 'E58A45', r=6); s.bulat(x + 2, y + 1.5, 9, 'C05A0B'); s.teks(x + 2, y + 1.5, 9, 9, str(n), 5.6, 'FFFFFF', True, 'c')
    s.teks(x + 14, y, w - 16, 12, [(t, NAVY), (f'{c} titik', ORG)], 6.0, tebal=True); x += w + 5
s.teks(46, 170, 200, 10, ('Titik panas 24 jam s.d. cut off' if HARI == CUT.date().isoformat() else f'Titik panas {HD}') + '; angka = 3 kecamatan teratas', 6.0, ABU)
ytl = 393; nk = sum(NLV.values())
s.teks(46, ytl, 120, 12, f"{f0(nk)} CLUSTER {int(D8[1][8:])}–{tg(HARI).upper()}", 6.8, '374151', True)
x = 160
for nm, c in (('Prioritas', MERAH), ('Siaga', 'E58A45'), ('Patroli', 'F4C3A0'), ('Rutin', 'A9B4C2')):
    s.bulat(x, ytl + 2.5, 7, c); tt = f'{nm} {NLV[nm]}'; s.teks(x + 10, ytl, 60, 12, tt, 6.8, MERAH if nm == 'Prioritas' else '374151', nm == 'Prioritas'); x += 14 + 3.7 * len(tt) + 8
# tabel kabupaten
R = [{'kab': k, 'hot': v['hot'][-1], 'luas': v['luas'][-1], 'belum': v['belum'][-1], 'k7': K7.get(k, 0)} for k, v in KAB.items()]
R.sort(key=lambda r: (-r['hot'], -r['luas'], -r['k7'], r['kab']))
s.teks(400, 167, 80, 16, 'KAB/KOTA', 6.4, ABU, True)
for x, w, t, c, ik in ((478, 44, 'TITIK*', NAVY, 'api'), (524, 52, 'LUAS (HA)*', SLATE, 'grid'), (578, 52, 'BELUM PADAM*', MERAH, 'awas'), (632, 55, 'CLUSTER 7 HARI', 'C05A0B', 'pin')):
    s.kotak(x, 167, w, 16, c, r=4); s.ikon(ik, 'putih', x + 3.5, 171, 7.5); s.teks(x + 12.5, 167, w - 13, 16, t, 4.6, 'FFFFFF', True, wrap=False)
COLX = [(500, 'hot'), (550, 'luas'), (604, 'belum'), (660, 'k7')]
for n, r in enumerate(R[:10]):
    yy = 188 + 18.4 * n
    if n == 0: s.kotak(392, yy - 2, 295, 18.4, 'FBEFE4')
    s.bulat(395, yy + 3, 11, NAVY); s.teks(395, yy + 3, 11, 11, str(n + 1), 5.8, 'FFFFFF', True, 'c')
    s.teks(411, yy, 70, 14.4, r['kab'].replace('Tanjung Jabung', 'Tanjab'), 7.4, '111827', True)
    for cx, kk in COLX: s.teks(cx - 25, yy, 50, 14.4, f0(r[kk]), 7.4, '111827', kk == 'hot', 'c')
    if n < 9: s.garis(392, yy + 16.2, 687, yy + 16.2, 'EEF1F5', .5)
s.kotak(392, 374, 295, 22, NAVY, r=4)
s.teks(400, 374, 80, 22, 'Jambi', 8, 'FFFFFF', True)
for cx, kk in COLX: s.teks(cx - 25, 374, 50, 22, f0(sum(r[kk] for r in R)), 8, 'FFFFFF', True, 'c')
s.teks(393, 398, 290, 10, f"*{CUT.day} Okt {CUT.year}, s.d. {JAM} WIB" if SEM else f'*{HD} {CUT.year}, hari penuh', 6, ABU)
# kejadian dan status
s.panel(696, 165, 221, 244, 'Kejadian dan status', 'awas', MERAH, 'F8DDE0', '7', MERAH)
kab_top = R[0]['kab']; kab2 = R[1]['kab']
KJ = []
for kk, nm in (('tn', 'TN Berbak, Tanjab Timur'), ('thr', 'Tahura OKH, Muaro Jambi')):
    v = KW[kk]; bl = v['belum'][-1]
    KJ.append(('api', ORG, nm, [(f"{f0(v['luas'][-1])} ha terbakar {HD}", 'n', 'grid'), (f'belum padam {f0(bl)} ha', 'm', 'awas')], ('belum padam', 'm', 'api') if bl >= 1 else (('pantau', 'n', 'mata') if SEM else ('padam', 'n', 'cek'))))
KJ.append(('api', MERAH, f"Kompleks {TOPK['id']}, {TOPK['kec']}", [(f"{f0(TOPK['luas'])} ha kompleks", 'o', 'grid'), (f"{f0(TOPK['n24'])} titik 24 jam", 'n', 'api')], (NAMA_LV[TOPK['level']], 'm' if TOPK['level'] == 3 else 'o', 'api')))
for kb in (kab_top, kab2):
    KJ.append(('mata', SLATE, kb, [(f"{f0(KAB[kb]['hot'][-1])} titik {HD}", 'o', 'api'), (f'{K7.get(kb, 0)} cluster', 'n', 'pin')], ('pantau', 'n', 'mata')))
KJ.append(('orang', NAVY, 'Personel lapangan Jambi', [(f"{f0(ADA['personel'])} personel", 'n', 'orang'), (f"{OPP['daops']['n']} Daops, {OPP['daops']['regu']} regu", 'n', 'perisai')], ('berjalan', 'o', 'putar')))
pct_rerata = re.search(r'(\d+)% dari rerata', D['kumulatif']['kalimat'] or '')
KJ.append(('lapis', NAVY, f'Karhutla {CUT.year}', [(f"{f0(D['luas_tahun'])} ha terbakar", 'o', 'grid'), (f"1 Jan–{tg(D['luas_tahun_sd'])}", 'n', 'kalender')] + ([(f'titik {pct_rerata.group(1)}% rerata', 'n', 'api')] if pct_rerata else []), ('kumulatif', 'n', 'lapis')))
for n, (ik, c, nm, ch, st) in enumerate(KJ):
    yy = 209 + 28.4 * n
    s.bulat(702, yy + 2, 16, c); s.ikon(ik, 'putih', 706, yy + 6, 8)
    s.teks(723, yy, 120, 11, nm, 7.5, '111827', True)
    w = 3.25 * len(st[0]) + 19; s.pil(908 - w, yy, st[0], CHIP[st[1]][0], CHIP[st[1]][1], st[2])
    x = 723
    for t, j, ii in ch:
        if x > 870: break
        x += s.pil(x, yy + 13, t, CHIP[j][0], CHIP[j][1], ii, sz=5.6, k=3.15) + 4
    if n < len(KJ) - 1: s.garis(702, yy + 26.5, 908, yy + 26.5, 'EEF1F5', .5)
# tren
s.kotak(43, 417, 874, 68, 'F3F5F8', r=8); s.label(52, 421, 'TREN TITIK PANAS')
tt, cc = pct(h[-2], h[-1])
s.teks(560, 421, 350, 12, [(tt + '  ', cc, True), (f"{HD_S} dibanding {tg(D8[-2])}", '4B5563', False)], 6.8, rata='r')
batang(s, h, 52, 856, 462, 24, 62, TGL, label_akhir=f'{TGL[-1]} Okt, sementara' if SEM else None, sz_v=7.6, sz_t=5.8)
s.catatan(CATATAN + f" Angka {HD} sementara, data s.d. {JAM} WIB." * SEM + f' Personel: rekap BNPB harian {tg(hari_ada)}; PM2.5: CAMS (perkiraan, belum divalidasi).')

# =============================================================== slide 2: kabupaten teratas
KB = kab_top; V = KAB[KB]; KBS = KB.replace('Tanjung Jabung', 'Tanjab')
s = Sl(2)
s.kepala('Sebaran Titik Panas dan Kebakaran di ', f'Kab. {KB}', f" per {int(HARI[8:])} {BLN_P[int(HARI[5:7]) - 1]}")
s.kotak(43, 81, 565, 78, 'EEF1F5', r=8); s.label(52, 87, 'DATA SATELIT')
s.kartu(52, 107, 177, 43, f0(V['hot'][-1]), f'titik panas {HD}', 'api', 'oranye', 'FBE7D6', ORG)
s.kartu(237, 107, 177, 43, f0(V['luas'][-1]) + ' ha', f'luas terbakar {HD}', 'grid', 'navy', 'E6EAF0', NAVY)
s.kartu(422, 107, 178, 43, f0(V['belum'][-1]) + ' ha', f'belum padam {HD}', 'awas', 'merah', 'F8DDE0', MERAH)
kk_kab = {k: v for k, v in KK.items() if k.startswith(KB + '|')}
kb_bel = sum(v['belum'][-1] for v in kk_kab.values()); kb_hot = sum(v['hot'][-1] for v in kk_kab.values())
s.kotak(630, 81, 287, 78, 'FBEDEF', r=8); s.label(639, 87, 'PERLU PERHATIAN', MERAH, 'awas')
s.kartu(638, 107, 132, 43, f0(kb_bel) + ' ha', f'belum padam di kawasan konservasi, {HD}', 'api', 'merah', 'F8DDE0', MERAH, sz=15, ket_sz=6.2)
s.kartu(777, 107, 132, 43, f0(kb_hot), f'titik di kawasan konservasi, {HD}', 'pohon', 'oranye', 'FBE7D6', ORG, sz=15, ket_sz=6.2)
s.panel(43, 165, 426, 320, 'Tren Harian', 'kalender', NAVY, 'EEF1F5')
for j, (m, judul) in enumerate((('hot', 'Titik panas'), ('luas', 'Luas terbakar (ha)'), ('belum', 'Belum padam (ha)'))):
    y0 = 210 + 92 * j; s.teks(54, y0, 150, 12, judul, 8.4, NAVY, True)
    batang(s, V[m], 54, 405, y0 + 56, 26, 31, TGL, label_akhir=f'{TGL[-1]} Okt, sementara' if SEM else None, sz_v=7.2, sz_t=5.6)
# per kecamatan
KC = sorted(((k.split('|')[1], sum(v['hot'])) for k, v in D['kec'].items() if k.startswith(KB + '|')), key=lambda r: -r[1])
n8 = sum(V['hot'])
s.panel(480, 165, 437, 182, 'Titik Panas per Kecamatan', 'gunung', ORG, 'FBE7D6', f0(n8), ORG)
mx = KC[0][1] if KC else 1
for n, (kec, c) in enumerate(KC[:6]):
    yy = 212 + 16.8 * n; inkaw = (KB + '|' + kec) in KK
    s.bulat(489, yy + 1.5, 10, NAVY); s.teks(489, yy + 1.5, 10, 10, str(n + 1), 5.6, 'FFFFFF', True, 'c')
    s.teks(504, yy, 90, 13, kec, 7.4, '111827', True)
    if inkaw: s.pil(593, yy + 1, 'kawasan', 'FBE7D6', ORG, 'pohon', sz=5.6)
    bw = max(3, 180 * c / mx); s.kotak(681, yy + 2.5, bw, 8.5, ORG if inkaw else 'A9B4C2', r=2)
    s.teks(681 + bw + 4, yy, 40, 13, f0(c), 7.4, '111827', True)
s.teks(489, 314, 300, 9, f"SEBARAN TITIK {int(D8[0][8:])}–{tg(HARI).upper()}", 6, '374151', True)
dk = sum(sum(v['hot']) for v in kk_kab.values()); lk = n8 - dk
w1 = max(170, min(420 - 120, 420 * dk / max(1, n8)))
s.kotak(489, 324, w1, 15, ORG, r=7.5); s.ikon('pohon', 'putih', 495, 327.5, 8)
s.teks(506, 324, w1 - 10, 15, [(f'{f0(dk)} titik ({round(100 * dk / max(1, n8))}%)', 'FFFFFF', True), ('  di kawasan konservasi', 'FFFFFF', False)], 6.4)
s.kotak(489 + w1 + 2, 324, 420 - w1 - 2, 15, 'E6EAF0', r=7.5); s.ikon('pin', 'navy', 489 + w1 + 8, 327.5, 8)
s.teks(489 + w1 + 19, 324, 420 - w1 - 20, 15, [(f'{f0(lk)} titik ({round(100 * lk / max(1, n8))}%)', NAVY, True), ('  di luar', NAVY, False)], 6.4)
# situasi lapangan
KX = max((k for k in JK if k['kab'] == KB), key=lambda k: (k['level'], k['luas']))
OX = next((c for c in OPK if c['id'] == KX['id']), None) or next((c for c in OPK if c['kab'] == KB), None)
KE = KEND.get(KX['id']) or {}
s.panel(480, 356, 437, 129, 'Situasi lapangan', 'awas', MERAH, 'F8DDE0', '3', MERAH)
SL = [('api', ORG, f"Kompleks {KX['id']} {KX['pj']}, Kec. {KX['kec']}",
       [(f"{f0(KX['luas'])} ha kompleks", 'm', 'grid'), ('gambut' if KX['gambut'] else 'tanah mineral', 'n', 'gunung')] + ([(f"{KE['air_jenis']} ±{f1(KE['air_km'])} km", 'n', 'air')] if KE.get('air_km') else []), (NAMA_LV[KX['level']], 'm', 'api')),
      ('rute', NAVY, 'Akses darat ke titik api', [(f"jalan utama ±{f1(KE.get('jalan_km', 0))} km", 'n', 'rute')] + ([(f"tiba ±{f1(OX['jam'])} jam dari {OX['dari'].split(' ', 1)[-1] if OX['dari'].startswith(('A', 'T')) and OX['dari'][1:4].isdigit() else OX['dari']}", 'n', 'jam')] if OX and OX.get('jam') else []), ('kendala', 'o', 'awas')),
      ('asap', SLATE, f"Udara {pm_lok if pm_kab else 'Jambi'}", [(f'PM2.5 {f0(pm_v)} µg/m³', 'n', 'asap'), (kat_pm(pm_v), 'm' if pm_v > PM['batas'][1] else 'n', 'awas'), (f'prakiraan CAMS {PMH}', 'g', 'mata')], ('pantau', 'n', 'mata'))]
for n, (ik, c, nm, ch, st) in enumerate(SL):
    yy = 402 + 27.5 * n
    s.bulat(489, yy + 2, 16, c); s.ikon(ik, 'putih', 493, yy + 6, 8)
    s.teks(511, yy, 260, 11, nm, 7.6, '111827', True)
    w = 3.25 * len(st[0]) + 19; s.pil(908 - w, yy, st[0], CHIP[st[1]][0], CHIP[st[1]][1], st[2])
    x = 511
    for t, j, ii in ch: x += s.pil(x, yy + 13, t, CHIP[j][0], CHIP[j][1], ii, sz=5.6, k=3.15) + 4
    if n < 2: s.garis(489, yy + 25.5, 908, yy + 25.5, 'EEF1F5', .5)
s.catatan(CATATAN + f" Angka {HD} sementara, data s.d. {JAM} WIB." * SEM + ' Kawasan konservasi = TN Berbak dan Tahura Orang Kayo Hitam; akses dan jarak air: lapisan jalan dan kanal aplikasi.')

# =============================================================== slide 3: Kawasan Berbak
s = Sl(3)
s.kepala('Sebaran Titik Panas dan Kebakaran di ', 'Kawasan Berbak', f" per {int(HARI[8:])} {BLN_P[int(HARI[5:7]) - 1]}")
A = KW['all']
s.kotak(43, 81, 565, 78, 'EEF1F5', r=8); s.label(52, 87, 'DATA SATELIT')
s.kartu(52, 107, 177, 43, f0(A['hot'][-1]), f'titik panas {HD}', 'api', 'oranye', 'FBE7D6', ORG)
s.kartu(237, 107, 177, 43, f0(A['luas'][-1]) + ' ha', f'luas terbakar {HD}', 'grid', 'navy', 'E6EAF0', NAVY)
s.kartu(422, 107, 178, 43, f0(A['belum'][-1]) + ' ha', f'belum padam {HD}', 'awas', 'merah', 'F8DDE0', MERAH)
EMP = [k for k in sorted(KK, key=lambda k: (-KK[k]['hot'][-1], -sum(KK[k]['hot']))) if KK[k]['hot'][-1] > 0][:2] or sorted(KK, key=lambda k: -sum(KK[k]['hot']))[:1]
top2 = [k.split('|')[1] for k in EMP]
s.kotak(630, 81, 287, 78, 'FBEDEF', r=8); s.label(639, 87, 'PERLU PERHATIAN', MERAH, 'awas')
s.kotak(638, 107, 271, 43, 'FFFFFF', 'D5DBE3', r=6); s.bulat(647, 116.5, 24, 'F8DDE0'); s.ikon('api', 'merah', 652, 121.5, 14)
s.teks(679, 109, 225, 20, ' dan '.join(top2), 15, MERAH, True, v='b')
nsg = sum(1 for k in KAW_JK if k['level'] >= 2)
s.teks(679, 129.5, 225, 19, f"{' dan '.join(f0(KK[k]['hot'][-1]) for k in EMP)} titik {HD_S}; belum padam kawasan {f0(A['belum'][-1])} ha; {nsg} kompleks siaga/prioritas di kawasan", 6.4, '374151', v='t')
# peta
s.panel(43, 165, 303, 320, 'Sebaran Titik Panas', 'peta', NAVY, 'EEF1F5')
P3 = json.load(open('peta3.json')); mw = 278; mh = mw * P3['h'] / P3['w']
s.gambar('peta3.png', 55.5, 210, mw, mh)
x = 58; per = [D8[0:3], D8[3:6], D8[6:8]]
for j, (c, t) in enumerate((('A9B4C2', f"{int(per[0][0][8:])}–{tg(per[0][-1])}"), ('F0A35E', f"{int(per[1][0][8:])}–{tg(per[1][-1])}"), (MERAH, f"{int(per[2][0][8:])}–{tg(per[2][-1])}, terbaru"))):
    s.bulat(x, 469, 7, c); s.teks(x + 10, 465, 80, 14, t, 6.6, MERAH if j == 2 else '374151', j == 2); x += 10 + 3.6 * len(t) + 14
# per kecamatan (maks 4); slot kosong diisi belum padam seluruh kawasan
s.panel(357, 165, 320, 320, 'Titik panas harian per kecamatan', 'api', ORG, 'FBE7D6')
CH = [(k.split('|')[1], 'titik di dalam kawasan', KK[k]['hot']) for k in sorted(KK, key=lambda k: -sum(KK[k]['hot']))[:4]]
if len(CH) < 4: CH.append(('Seluruh kawasan', 'belum padam (ha)', A['belum']))
if len(CH) < 4: CH.append(('Seluruh kawasan', 'luas terbakar (ha)', A['luas']))
for n, (nm, sub, v) in enumerate(CH[:4]):
    cx0 = 366 + 152 * (n % 2); cy0 = 212 + 110 * (n // 2)
    s.teks(cx0, cy0, 140, 12, nm, 8.4, NAVY, True); s.teks(cx0, cy0 + 11, 140, 9, sub, 5.8, ABU)
    batang(s, v, cx0, 145, cy0 + 82, 44, 11.5, TGL, pcts=False, sz_v=6.2, sz_t=5.4)
# alur api: rentang hari aktif tiap kecamatan (hari bersebelahan >= 25% puncaknya)
alur = []
for k in sorted(KK, key=lambda k: -sum(KK[k]['hot']))[:4]:
    v = KK[k]['hot']; pk = max(range(8), key=lambda i: (v[i], i)); a = b2 = pk
    while a > 0 and v[a - 1] >= .25 * v[pk]: a -= 1
    while b2 < 7 and v[b2 + 1] >= .25 * v[pk]: b2 += 1
    alur.append((a, b2, k.split('|')[1]))
alur.sort()
s.teks(366, 432, 100, 9, 'ALUR API', 6, ABU, True)
x, y = 366, 444
for n, (a, b2, kec) in enumerate(alur):
    t = f"{kec} {TGL[a]}–{TGL[b2]}" if a != b2 else f"{kec} {TGL[a]}"
    j = 'm' if b2 >= 6 else 'o' if b2 >= 3 else 'n'
    w = 3.35 * len(t) + 19
    if x + w > 668: x, y = 366, y + 15
    s.pil(x, y, t, CHIP[j][0], CHIP[j][1], 'api', sz=5.8, h=12); x += w + 5
# kronologi dan unsur
s.panel(688, 165, 229, 320, 'Kronologi dan unsur', 'kalender', NAVY, 'EEF1F5')
awal = min((k for k in KAW_JK if k['awal']), key=lambda k: k['awal'], default=None)
pk = max(range(8), key=lambda i: A['hot'][i]); pk_kec = max(KK, key=lambda k: KK[k]['hot'][pk]).split('|')[1]
KR = []
if awal: KR.append(('api', ORG, tg(awal['awal']), [('api teramati', 'n', 'api'), (f"{awal['kec']}, {awal['id']}", 'n', 'pin')]))
KR.append(('aktivitas', ORG, tg(D8[pk]), [(f"puncak {f0(A['hot'][pk])} titik", 'o', 'api'), (pk_kec, 'n', 'pin')]))
KR.append(('air', NAVY, tg(D8[-2]), [(f"{f0(A['luas'][-2])} ha terbakar", 'o', 'grid'), (f"belum padam {f0(A['belum'][-2])} ha", 'm', 'awas')]))
KR.append(('awas', MERAH, HD_S, [(f"{f0(A['hot'][-1])} titik", 'm', 'api'), (f"belum padam {f0(A['belum'][-1])} ha", 'm', 'awas')]))
for n, (ik, c, jd, ch) in enumerate(KR):
    yy = 213 + 35 * n
    s.bulat(697, yy, 15, c, 'FFFFFF' if n < len(KR) - 1 else None); s.ikon(ik, 'putih', 700.5, yy + 3.5, 8)
    s.teks(717, yy - 1, 150, 11, jd, 7.6, '111827', True)
    x = 717
    for t, j, ii in ch: x += s.pil(x, yy + 11, t, CHIP[j][0], CHIP[j][1], ii, sz=5.6, k=3.15) + 4
pihak = []
for k in KAW_JK:
    if 'Taman Nasional' in k['pihak'] and 'Balai TN Berbak dan Sembilang' not in pihak: pihak.append('Balai TN Berbak dan Sembilang')
    if 'Pemerintah daerah' in k['pihak'] and 'Dishut Prov. Jambi (Tahura)' not in pihak: pihak.append('Dishut Prov. Jambi (Tahura)')
dop = Counter(k['daops'] for k in KAW_JK if k['daops']).most_common(1)
DAOPS = {x['kode']: x['nama'] for x in json.load(open('../../../public/data/daops.json')) if isinstance(x, dict) and 'kode' in x}
UN = [(p, 'pohon') for p in pihak] + ([(DAOPS.get(dop[0][0], f'Daops {dop[0][0]}'), 'perisai')] if dop else [])
NMU = {'Manggala Agni': 'Manggala Agni', 'MPA': 'MPA', 'TNI': 'TNI', 'Polri': 'Polri', 'personil BPBD': 'BPBD', 'Perusahaan': 'Perusahaan', 'Pemda': 'Pemda', 'Satpol PP': 'Satpol PP'}
IKU = {'Manggala Agni': 'api', 'MPA': 'orang', 'TNI': 'bendera', 'Polri': 'perisai', 'BPBD': 'perisai', 'Perusahaan': 'truk', 'Pemda': 'perisai', 'Satpol PP': 'perisai'}
for k, v in sorted(ADA['rinci'].items(), key=lambda kv: -kv[1]):
    if v > 0: UN.append((NMU[k], IKU[NMU[k]]))
UN = UN[:10]
s.teks(697, 357, 200, 10, f'UNSUR TERKAIT ({len(UN)})', 6.6, '374151', True)
for n, (t, ik) in enumerate(UN):
    cx = 697 + 107 * (n % 2); cy = 370 + 21.5 * (n // 2)
    s.kotak(cx, cy, 102, 16, 'EEF1F5', r=8); s.ikon(ik, 'navy', cx + 5, cy + 4, 8); s.teks(cx + 17, cy, 84, 16, t, 5.9 if len(t) > 22 else 6.3, NAVY, True)
s.catatan(CATATAN + ' Kawasan Berbak = TN Berbak dan Tahura Orang Kayo Hitam; unsur: pengelola kawasan dan rekap BNPB harian Jambi.')

# =============================================================== slide 4: kebutuhan pemadaman
s = Sl(4)
KE0 = KEND.get(OPK[0]['id']) or {}
air0 = KE0.get('air_km')
s.teks(43, 22, 820, 60, [('Kebutuhan pemadaman Kawasan Berbak: ', NAVY), (f"{f0(A['belum'][-1])} ha", ORG), (f" belum padam {HD}, kanal terdekat ±{f1(air0)} km dari kompleks terbesar" if air0 else f' belum padam {HD}', NAVY)], 21, tebal=True, v='t', spasi=1.0)
s.kepala('', None, '', cutoff=False, dua_baris=True)
s.kotak(43, 94, 565, 74, 'EAF1F9', r=8); s.label(52, 99, 'KONDISI LAPANGAN', SLATE, 'air')
K0 = next((k for k in JK if k['id'] == OPK[0]['id']), TOPK)
gb = D['gambut']
KL4 = [('Kering', f"DC {f0(K0['dc'])}, gambut sangat kering" if K0.get('dc') else 'musim kemarau', 'air', SLATE),
       ('Angin', f"{f1(K0['ws'])} km/jam, {TOPK['kec']}" if K0.get('ws') is not None else 'pantau BMKG', 'angin', NAVY),
       ('Gambut', f"{round(100 * gb[0] / max(1, gb[1]))}% titik {HD} di gambut", 'awas', ORG),
       (f'PM2.5 {f0(pm_v)}', f"{kat_pm(pm_v)}, {pm_short}, prakiraan {PMH}", 'asap', NAVY)]
for n, (a, b_, ik, c) in enumerate(KL4):
    x = 52 + 139 * n; s.kotak(x, 119, 131, 41, 'FFFFFF', 'D5DBE3', r=6)
    s.bulat(x + 7, 127.5, 24, 'FBE7D6' if c == ORG else 'E6EAF0'); s.ikon(ik, ik_warna(c), x + 12, 132.5, 14)
    s.teks(x + 37, 120, 92, 21, a, 15, c if c != SLATE else '5B7A99', True, v='b'); s.teks(x + 37, 142, 92, 16, b_, 6.4, '374151', v='t')
s.kotak(630, 94, 287, 74, 'FBEDEF', r=8); s.label(639, 99, 'KENDALA', MERAH, 'awas')
s.kotak(638, 119, 271, 41, 'FFFFFF', 'D5DBE3', r=6); s.bulat(645, 127.5, 24, 'F8DDE0'); s.ikon('air', 'merah', 650, 132.5, 14)
lbl = sorted(Counter(l for c in OPK for l in (KEND.get(c['id']) or {}).get('label', [])).items(), key=lambda kv: -kv[1])
judul_k = 'Jalan utama jauh dari titik api' if lbl and 'Jalan' in lbl[0][0] else (lbl[0][0] if lbl else 'Akses gambut')
s.teks(675, 120, 230, 18, judul_k, 10.5, MERAH, True, v='b')
x = 675
for t, ii in ((f"{lbl[0][1]} dari {len(OPK)} kompleks" if lbl else '', 'pin'), (f"jalan ±{f1(KE0.get('jalan_km', 0))} km di {OPK[0]['id']}", 'rute')):
    if t: x += s.pil(x, 143, t, 'F8DDE0', MERAH, ii, sz=5.6, k=3.15) + 5
# tabel kebutuhan: model operasi aplikasi, target terkurung 3 hari, skenario tengah; rentang = optimis–pesimis
B0, B1, B2 = BUT
s.panel(43, 175, 577, 229, 'Kebutuhan sarpras pemadaman', 'truk', NAVY, 'EEF1F5', None)
juta = lambda v: f1(v / 1e6) + ' juta L'
ITEM = [('POMPA DAN AIR', 'air', [
            ('Pompa pemadam', f"{f0(B1['pompa'])} unit, 2 per regu gambut ditambah cadangan", f"{f0(B0['pompa'])}–{f0(B2['pompa'])}"),
            ('Titik sumur bor', f"{f0(B1['bor'])} titik, 1 titik per {f0(D['koef']['ha_per_bor']['mid'])} ha bara gambut", f"{f0(B0['bor'])}–{f0(B2['bor'])}"),
            ('Air pembasahan bara', f"{juta(B1['air_l'])} untuk {f0(B1['bara'])} ha bara", f"{juta(B0['air_l'])}–{juta(B2['air_l'])}")]),
        ('SELANG', 'selang', [
            ('Selang penyalur', f"{f1(B1['selang'] / 1000)} km, {f0(D['koef']['selang_pompa']['mid'])} m per pompa", f"{f1(B0['selang'] / 1000)}–{f1(B2['selang'] / 1000)} km")]),
        ('PERSONEL', 'orang', [
            ('Regu garis api', f"{f0(B1['regu_garis'])} regu kerja untuk mengurung front", f"{f0(B0['regu_garis'])}–{f0(B2['regu_garis'])}"),
            ('Regu mop-up bara', f"{f0(B1['regu_bara'])} regu kerja, padam tuntas gambut", f"{f0(B0['regu_bara'])}–{f0(B2['regu_bara'])}"),
            ('Total personel', f"{f0(B1['personel'])} orang ({f0(B1['regu'])} regu dengan rotasi dan dukungan)", f"{f0(B0['personel'])}–{f0(B2['personel'])}")])]
nitem = sum(len(g[2]) for g in ITEM)
s.teks(560, 175, 50, 38, str(nitem), 15, NAVY, True, 'r')
s.teks(53, 219, 120, 10, 'KEBUTUHAN', 6.2, ABU, True); s.teks(200, 219, 150, 10, 'PERKIRAAN (TENGAH)', 6.2, ABU, True); s.teks(505, 219, 100, 10, 'RENTANG', 6.2, ABU, True)
y = 234; no = 1
for gj, gik, items in ITEM:
    s.kotak(52, y, 560, 12, 'F1F3F6', r=3); s.ikon(gik, 'navy', 56, y + 2.5, 7); s.teks(67, y, 120, 12, gj, 6.3, NAVY, True)
    s.pil(67 + 4.6 * len(gj) + 10, y + 1.5, f'{len(items)} butir', 'E1E6ED', NAVY, 'lapis', sz=5, h=9, k=2.9)
    y += 15.5
    for nm, rinci, rentang in items:
        s.bulat(56, y + 1.5, 9, NAVY); s.teks(56, y + 1.5, 9, 9, str(no), 5.2, 'FFFFFF', True, 'c')
        s.teks(70, y, 125, 12, nm, 7.2, '111827', True); s.teks(200, y, 300, 12, rinci, 7.2, '111827', True)
        s.pil(505, y + .5, rentang, 'E6EAF0', SLATE, 'aktivitas', sz=5.6, k=3.1)
        y += 15.2; no += 1
s.teks(53, 389, 560, 10, [('DASAR:  ', ABU, True), (f"model operasi aplikasi untuk {len(OPK)} kompleks api di kawasan, target api terkurung 3 hari; rentang = skenario optimis–pesimis.", '4B5563', False)], 6)
# sarana dan personel
s.panel(630, 175, 287, 229, 'Sarana dan personel Jambi', 'truk', NAVY, 'EEF1F5')
rc = ADA['rinci']
SR = [(f0(ADA['personel']), f'personel, rekap BNPB {tg(hari_ada)}', 'orang'), (f0(OPP['daops']['n']), f"Daops Manggala Agni, {OPP['daops']['regu']} regu", 'perisai'),
      (f0(rc.get('MPA', 0)), 'Masyarakat Peduli Api', 'orang'), (f0(rc.get('Perusahaan', 0)), 'personel perusahaan', 'truk'),
      (f0(rc.get('TNI', 0) + rc.get('Polri', 0)), 'TNI dan Polri', 'bendera'), (f0(ADA['wb']), f'heli water bombing, rekap {tg(hari_ada)}', 'pesawat')]
for n, (a, b_, ik) in enumerate(SR):
    x = 639 + 137 * (n % 2); y = 222 + 58 * (n // 2)
    s.kotak(x, y, 131, 50, 'F3F5F8', r=6); s.bulat(x + 8, y + 13, 24, 'E6EAF0'); s.ikon(ik, 'slate', x + 13, y + 18, 14)
    s.teks(x + 40, y + 7, 88, 20, a, 15, NAVY, True, v='b'); s.teks(x + 40, y + 28, 88, 18, b_, 6.4, '374151', v='t')
# rencana tindak lanjut (usulan dari data)
s.kotak(43, 413, 874, 72, 'FBF1E8', r=8); s.teks(55, 418, 200, 11, 'RENCANA TINDAK LANJUT (USULAN)', 7, 'C05A0B', True)
OK0 = OPK[0]
RTL = [('bor', f"Sumur bor di {OK0['kec']}", f"Buat {f0(OK0['bor'][1])} titik bor di kompleks {OK0['id']} ({OK0['pj']}); kanal terdekat ±{f1(KE0.get('air_km', 0))} km"),
       ('orang', 'Regu mop-up gambut', f"Siapkan {f0(B1['regu_bara'])} regu bara untuk {f0(B1['bara'])} ha bara gambut di kawasan, target 3 hari"),
       ('rute', 'Akses ke titik api', f"Jalur darat dari {OK0['dari'].split(' ', 1)[-1] if OK0['dari'][1:4].isdigit() else OK0['dari']}, tiba ±{f1(OK0['jam'])} jam; jalan utama ±{f1(KE0.get('jalan_km', 0))} km" if OK0.get('jam') else 'Buka jalur darat dan sungai ke titik api'),
       ('asap', 'Lindungi warga dari asap', f"Prakiraan PM2.5 {pm_short} {PMH} {f0(pm_v)} µg/m³ ({kat_pm(pm_v)}); siapkan masker dan posko kesehatan")]
for n, (ik, a, b_) in enumerate(RTL):
    x = 55 + 216 * n
    s.bulat(x, 438, 26, ORG); s.ikon(ik, 'putih', x + 6, 444, 14)
    s.teks(x + 33, 432, 175, 12, a, 8.2, NAVY, True); s.teks(x + 33, 445, 172, 34, b_, 6.6, '4B5563', v='t')
    if n: s.garis(x - 9, 433, x - 9, 478, 'E5CDB8', .6)
s.catatan(CATATAN + f" Kebutuhan dan RTL: perkiraan model operasi aplikasi (koefisien terbuka), bukan perintah operasi; personel: rekap BNPB harian {tg(hari_ada)}.")

prs.core_properties.title = f"Paparan Karhutla Provinsi Jambi {CUT.strftime('%d%m%y')}"; prs.core_properties.author = 'BNPB'
prs.save(OUT)
bad = [(n, sh.text_frame.text) for n, sl in enumerate(prs.slides, 1) for sh in sl.shapes if sh.has_text_frame and re.search('[·—]|Lampung|Way Kambas', sh.text_frame.text)]
print('sisa terlarang', bad); print('simpan', OUT)
